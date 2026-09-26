"""Fixed-template python-pptx renderer and structural QA (CLAUDE.md section 9). No LLM, no template file.

`render(ctx)` runs the gate first (gate.run_gate) and refuses on FAIL; otherwise it builds outputs/<run_id>/pitch.pptx
from the audited deck and runs `structural_qa` on the saved file.

Template (all constants below; the LLM never controls layout): 16:9, a navy title bar (#002C77) with the fixed
slide title in white, one accent colour, Arial (its glyph for ₹ was checked), footer "Prepared by Marsh |
Confidential | <date>" and "n / 5", no logos.
- Only claims whose state isn't REMOVED and whose audit status isn't UNSUPPORTED or CONTRADICTED are rendered
  (ADVISOR_ATTESTED with its label). Company claims show the two user labels only (CLAUDE.md section 5):
  "(Web-sourced)" or "(Assumption)".
- A VERIFIED_WITH_QUALIFIER claim gets a numbered marker; its qualifier (Claim.qualifier_text) is a small footnote at
  the bottom of that slide and is collected again on slide 5.
- Slide 1: company name, company claims, the exposures considered (code). Slide 2: the WM claims. Slide 3: a real
  table Exposure | Benefit | Condition / limitation | Source. Slide 4: the recommended policy block (code-injected
  name, variant, add-ons, decided by), the reasons with their company framing as smaller sub-lines, and a column
  of supporting benefits and key limitations. Slide 5: qualifiers and key terms, advisor-attested claims,
  assumptions, sources, disclaimer.
- Speaker notes on every slide: claim_id → status → evidence_id (document, page) for each claim of that slide.
- Text never overflows: each box's font size is the largest that fits, estimated with the real Arial glyph widths
  (Pillow) and word wrapping; `structural_qa` re-measures every box with the same estimator.

`structural_qa(path)`: 5 slides; titles in order; no "{{"; no backtick before a digit; no "Unverified" or raw fact
status; bullet counts within the
limits; exactly one policy named in slide 4's recommendation block; speaker notes on every slide; no text box or
table overflowing its space; file > 10 KB. Raises RenderQAError listing every problem.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from functools import lru_cache
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from marsh.audit import marsh_evidence
from marsh.decision_log import log_decision
from marsh.evidence_store import load_evidence, to_display
from marsh.gate import is_rendered, run_gate
from marsh.grounding import named_policies
from marsh.models import (
    SLIDE1_MAX_BULLETS,
    SLIDE2_MAX_BULLETS,
    SLIDE3_MAX_ROWS,
    SLIDE4_MAX_KEY_LIMITATIONS,
    SLIDE4_MAX_POLICY_BULLETS,
    SLIDE4_MAX_SUPPORTING_BENEFITS,
    SLIDE_TITLES,
    WEB_SOURCED_LABEL,
    AuditResult,
    AuditStatus,
    Claim,
    GateResult,
    OverallFlag,
    RunContext,
    save_json,
)
from marsh.pitch import ASSUMPTION_LABEL, NOT_SOURCE_VERIFIED_TEXT, NOT_STATED_TEXT, PITCH_FILE
from marsh.run_context import run_dir, save_run_context

PPTX_FILE = "pitch.pptx"
MIN_FILE_BYTES = 10 * 1024

# --- Template constants ------------------------------------------------------------------------------------------
SLIDE_W, SLIDE_H = 13.333, 7.5  # inches, 16:9
NAVY = RGBColor(0x00, 0x2C, 0x77)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x00, 0x9D, 0xE0)
TEXT = RGBColor(0x1E, 0x1E, 0x1E)
GREY = RGBColor(0x5F, 0x63, 0x68)
ROW_SHADE = RGBColor(0xEE, 0xF2, 0xF8)
FONT = "Arial"  # renders ₹ (U+20B9): its glyph differs from the missing-glyph box (checked with Pillow)
MARGIN = 0.5
TITLE_BAR_H = 0.95
ACCENT_H = 0.06
BODY_TOP = 1.25
FOOTER_Y = 7.08
BODY_BOTTOM = 6.95  # footnotes and body end above the footer
INSET = 0.08  # text box margins
LINE_SPACING = 1.2  # line height / font size (single spacing)
FOOTNOTE_PT = 9


class RenderQAError(RuntimeError):
    def __init__(self, problems: list[str]):
        super().__init__("the rendered deck failed structural QA: " + "; ".join(problems))
        self.problems = problems


class RenderRefusedError(RuntimeError):
    """The gate says FAIL: nothing is rendered."""

    def __init__(self, gate: GateResult):
        super().__init__("the gate FAILed: " + "; ".join(f.message for f in gate.failures))
        self.gate = gate


# --- Text measurement (renderer and QA) --------------------------------------------------------------------------

_FONT_FILES = {(False, False): "arial.ttf", (True, False): "arialbd.ttf", (False, True): "ariali.ttf",
               (True, True): "arialbi.ttf"}
_NARROW = set("iljtfrI.,;:'!|()[] ")
_WIDE = set("mwMW@%")


@lru_cache(maxsize=16)
def _pil_font(bold: bool, italic: bool):
    try:
        from PIL import ImageFont

        return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / _FONT_FILES[(bold, italic)]), 100) \
            if Path(r"C:\Windows\Fonts").exists() else ImageFont.truetype(_FONT_FILES[(bold, italic)], 100)
    except (OSError, ImportError):
        return None


def text_width_pt(text: str, size_pt: float, bold: bool = False, italic: bool = False) -> float:
    """Width of a single line of Arial text in points (Pillow glyph metrics; a character table without Pillow)."""
    font = _pil_font(bold, italic)
    if font is not None:
        return font.getlength(text) * size_pt / 100
    em = sum(0.28 if c in _NARROW else 0.83 if c in _WIDE else 0.67 if c.isupper() else 0.56 for c in text)
    return em * size_pt * (1.05 if bold else 1.0)


def wrapped_lines(text: str, size_pt: float, width_in: float, bold: bool = False, italic: bool = False) -> int:
    """Lines a paragraph needs at this size and width (greedy word wrap, as PowerPoint wraps)."""
    width = width_in * 72
    lines = 0
    for part in (text or " ").split("\n"):
        line, count = "", 1
        for word in part.split(" "):
            candidate = f"{line} {word}" if line else word
            if line and text_width_pt(candidate, size_pt, bold, italic) > width:
                count += 1
                line = word
            else:
                line = candidate
        lines += count
    return lines


@dataclass
class Run:
    text: str
    size: float
    bold: bool = False
    italic: bool = False
    color: RGBColor = TEXT
    superscript: bool = False


@dataclass
class Para:
    runs: list[Run]
    space_before: float = 0.0  # points
    indent: float = 0.0  # inches (left margin of the text, bullet hangs in it)
    bullet: str | None = None
    align: PP_ALIGN = PP_ALIGN.LEFT

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)


def para_height(text: str, size: float, bold: bool, italic: bool, space_before: float, indent: float,
                width_in: float) -> float:
    """Height in inches of one paragraph."""
    lines = wrapped_lines(text, size, max(0.2, width_in - indent), bold, italic)
    return (lines * size * LINE_SPACING + space_before) / 72


def paras_height(paras: list[Para], width_in: float) -> float:
    total = 0.0
    for p in paras:
        size = max(r.size for r in p.runs)
        total += para_height(p.text, size, any(r.bold for r in p.runs), all(r.italic for r in p.runs),
                             p.space_before, p.indent, width_in - 2 * INSET)
    return total + 2 * INSET


# --- Drawing ------------------------------------------------------------------------------------------------------


def _emu(inches: float) -> Emu:
    return Emu(int(round(inches * 914400)))


def _set_bullet(paragraph, char: str, indent_in: float) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(_emu(indent_in)))
    p_pr.set("indent", str(-_emu(min(indent_in, 0.22))))
    for tag in ("a:buNone", "a:buChar", "a:buFont", "a:buClr"):
        for old in p_pr.findall(qn(tag)):
            p_pr.remove(old)
    bu_clr = etree.SubElement(p_pr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", str(ACCENT))
    etree.SubElement(p_pr, qn("a:buFont")).set("typeface", FONT)
    etree.SubElement(p_pr, qn("a:buChar")).set("char", char)


def _indent_only(paragraph, indent_in: float) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(_emu(indent_in)))
    p_pr.set("indent", "0")
    etree.SubElement(p_pr, qn("a:buNone"))


def _fill_frame(frame, paras: list[Para], anchor=MSO_ANCHOR.TOP) -> None:
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, _emu(INSET))
    for n, para in enumerate(paras):
        p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
        p.alignment = para.align
        p.line_spacing = 1.0
        if para.space_before:
            p.space_before = Pt(para.space_before)
        for run in para.runs:
            r = p.add_run()
            r.text = run.text
            font = r.font
            font.name, font.size, font.bold, font.italic = FONT, Pt(run.size), run.bold, run.italic
            font.color.rgb = run.color
            if run.superscript:
                r._r.get_or_add_rPr().set("baseline", "30000")
        if para.bullet:
            _set_bullet(p, para.bullet, para.indent)
        elif para.indent:
            _indent_only(p, para.indent)


def _textbox(slide, name: str, x: float, y: float, w: float, h: float, paras: list[Para],
             anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    box.name = name
    _fill_frame(box.text_frame, paras, anchor)
    return box


def _fit(build, width: float, height: float, max_size: float, min_size: float) -> tuple[list[Para], float]:
    """The paragraphs at the largest base size (step 0.5 pt) whose estimated height fits the box."""
    size = max_size
    while size > min_size and paras_height(build(size), width) > height:
        size -= 0.5
    return build(size), size


def _rect(slide, x, y, w, h, color: RGBColor, name: str):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(x), _emu(y), _emu(w), _emu(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


# --- Deck model -> slides -----------------------------------------------------------------------------------------


@dataclass
class _Deck:
    """What the renderer needs from the run: rendered claims, results, footnote numbers, evidence lookups."""
    ctx: RunContext
    results: dict[str, AuditResult]
    notes_lookup: dict[str, str] = field(default_factory=dict)
    footnotes: dict[str, int] = field(default_factory=dict)  # claim_id -> qualifier number
    qualifier_texts: list[str] = field(default_factory=list)

    def shown(self, claim: Claim | None) -> bool:
        return claim is not None and is_rendered(claim, self.results.get(claim.claim_id))

    def status(self, claim: Claim) -> AuditStatus | None:
        result = self.results.get(claim.claim_id)
        return result.status if result else None

    def marker(self, claim: Claim) -> str | None:
        """The claim's qualifier number (assigned in slide order), if it has a rendered qualifier."""
        result = self.results.get(claim.claim_id)
        if not (result and result.status == AuditStatus.VERIFIED_WITH_QUALIFIER and claim.qualifier_text):
            return None
        if claim.claim_id not in self.footnotes:
            text = to_display(claim.qualifier_text)
            if text not in self.qualifier_texts:  # the same qualifier keeps one number
                self.qualifier_texts.append(text)
            self.footnotes[claim.claim_id] = self.qualifier_texts.index(text) + 1
        return str(self.footnotes[claim.claim_id])


# Code-made fixed lines of decks generated before the two-label rule (CLAUDE.md section 5) → today's wording.
LEGACY_FIXED_TEXT = {"Company details are AI-generated from model knowledge and unverified.": NOT_SOURCE_VERIFIED_TEXT}


def claim_text(claim: Claim, deck: _Deck) -> str:
    """The claim as shown: ₹ display, and the user labels (company claims, advisor attestation)."""
    text = to_display(claim.text)
    if claim.claim_type.value == "NON_FACTUAL":
        text = LEGACY_FIXED_TEXT.get(text, text)
    if claim.policy_id is None and claim.slide_number in (1, 4) and claim.basis_fact_ids:
        if claim.qualifier_text == WEB_SOURCED_LABEL:
            text = text.removesuffix(ASSUMPTION_LABEL.strip()).rstrip() + f" ({WEB_SOURCED_LABEL})"
        elif not text.rstrip().endswith(ASSUMPTION_LABEL.strip()):
            text = text.rstrip() + ASSUMPTION_LABEL  # every other company claim is an Assumption to the user
    if deck.status(claim) == AuditStatus.ADVISOR_ATTESTED:
        text += " (Advisor-attested)"
    return text


def _claim_runs(claim: Claim, deck: _Deck, size: float, **style) -> list[Run]:
    runs = [Run(claim_text(claim, deck), size, **style)]
    if (n := deck.marker(claim)) is not None:
        runs.append(Run(f"[{n}]", size, bold=True, color=ACCENT, superscript=True))
    return runs


def _footnote_paras(numbers: list[int], deck: _Deck) -> list[Para]:
    return [Para([Run(f"[{n}] {deck.qualifier_texts[n - 1]}", FOOTNOTE_PT, color=GREY)]) for n in numbers]


def _slide_footnotes(slide, claims: list[Claim], deck: _Deck) -> float:
    """Draw the slide's qualifier footnotes above the footer; returns their top (the body's bottom limit)."""
    numbers = sorted({deck.footnotes[c.claim_id] for c in claims if c.claim_id in deck.footnotes})
    if not numbers:
        return BODY_BOTTOM
    paras = _footnote_paras(numbers, deck)
    width = SLIDE_W - 2 * MARGIN
    height = paras_height(paras, width)
    top = BODY_BOTTOM - height
    _textbox(slide, "Footnotes", MARGIN, top, width, height, paras)
    return top - 0.05


def _base_slide(prs, number: int, date: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _rect(slide, 0, 0, SLIDE_W, TITLE_BAR_H, NAVY, "TitleBar")
    _rect(slide, 0, TITLE_BAR_H, SLIDE_W, ACCENT_H, ACCENT, "AccentLine")
    _textbox(slide, "Title", MARGIN, 0.12, SLIDE_W - 2 * MARGIN, TITLE_BAR_H - 0.2,
             [Para([Run(SLIDE_TITLES[number - 1], 28, bold=True, color=WHITE)])], anchor=MSO_ANCHOR.MIDDLE)
    _textbox(slide, "Footer", MARGIN, FOOTER_Y, 8.5, 0.32,
             [Para([Run(f"Prepared by Marsh | Confidential | {date}", 9, color=GREY)])])
    _textbox(slide, "SlideNumber", SLIDE_W - MARGIN - 1.2, FOOTER_Y, 1.2, 0.32,
             [Para([Run(f"{number} / 5", 9, color=GREY)], align=PP_ALIGN.RIGHT)])
    return slide


def _bullet_paras(claims: list[Claim], deck: _Deck, size: float, gap: float) -> list[Para]:
    return [Para(_claim_runs(c, deck, size), space_before=gap if n else 0, indent=0.3, bullet="•")
            for n, c in enumerate(claims)]


def _slide1(prs, deck: _Deck, date: str):
    ctx = deck.ctx
    slide = _base_slide(prs, 1, date)
    claims = [c for c in ctx.deck.slides[0].bullets if deck.shown(c)]
    for c in claims:
        deck.marker(c)
    bottom = _slide_footnotes(slide, claims, deck)
    width = SLIDE_W - 2 * MARGIN
    _textbox(slide, "CompanyName", MARGIN, BODY_TOP, width, 0.55,
             [Para([Run(ctx.company_name, 24, bold=True, color=NAVY)])])
    considered = [e.name for e in ctx.exposures if not e.assumption_based]
    assumed = [e.name for e in ctx.exposures if e.assumption_based]
    exposure_text = "Employee-health exposures considered: " + (", ".join(considered) or "none")
    if assumed:
        exposure_text += "; based on assumptions: " + ", ".join(assumed)
    exp_paras = [Para([Run(exposure_text, 12, italic=True, color=GREY)])]
    exp_h = paras_height(exp_paras, width)
    _textbox(slide, "Exposures", MARGIN, bottom - exp_h, width, exp_h, exp_paras)
    top = BODY_TOP + 0.65
    paras, _ = _fit(lambda s: _bullet_paras(claims, deck, s, s * 0.6), width, bottom - exp_h - 0.1 - top, 20, 11)
    _textbox(slide, "Body", MARGIN, top, width, bottom - exp_h - 0.1 - top, paras)
    return slide, claims


def _slide2(prs, deck: _Deck, date: str):
    slide = _base_slide(prs, 2, date)
    claims = [c for c in deck.ctx.deck.slides[1].bullets if deck.shown(c)]
    for c in claims:
        deck.marker(c)
    bottom = _slide_footnotes(slide, claims, deck)
    width = SLIDE_W - 2 * MARGIN
    paras, _ = _fit(lambda s: _bullet_paras(claims, deck, s, s * 0.9), width, bottom - BODY_TOP - 0.1, 22, 12)
    _textbox(slide, "Body", MARGIN, BODY_TOP + 0.1, width, bottom - BODY_TOP - 0.1, paras)
    return slide, claims


TABLE_COLS = (("Exposure", 2.25), ("Benefit", 4.45), ("Condition / limitation", 3.95), ("Source", 1.68))
CELL_INSET = 0.06


def _cell_paras(claim: Claim | None, deck: _Deck, size: float) -> list[Para]:
    if claim is None or not deck.shown(claim):
        return [Para([Run("—", size, color=GREY)])]
    if claim.text == NOT_STATED_TEXT:
        return [Para([Run(NOT_STATED_TEXT, size, italic=True, color=GREY)])]
    return [Para(_claim_runs(claim, deck, size))]


def _cell_height(paras: list[Para], width: float) -> float:
    return paras_height(paras, width - 2 * CELL_INSET + 2 * INSET) - 2 * INSET + 2 * CELL_INSET


def table_rows(deck: _Deck):
    return [r for r in deck.ctx.deck.slides[2].table_rows if deck.shown(r.benefit)][:SLIDE3_MAX_ROWS]


def _table_content(rows, deck: _Deck, size: float) -> list[list[list[Para]]]:
    header = [[Para([Run(name, size, bold=True, color=WHITE)])] for name, _ in TABLE_COLS]
    body = [[[Para([Run(r.exposure_name, size, bold=True, color=NAVY)])], _cell_paras(r.benefit, deck, size),
             _cell_paras(r.condition, deck, size), [Para([Run(to_display(r.source), size - 1.5, color=GREY)])]]
            for r in rows]
    return [header] + body


def _table_height(content, widths) -> list[float]:
    return [max(_cell_height(cell, w) for cell, w in zip(row, widths)) for row in content]


def _slide3(prs, deck: _Deck, date: str):
    slide = _base_slide(prs, 3, date)
    rows = table_rows(deck)
    claims = [c for r in rows for c in (r.benefit, r.condition) if deck.shown(c)]
    for c in claims:
        deck.marker(c)
    bottom = _slide_footnotes(slide, claims, deck)
    widths = [w for _, w in TABLE_COLS]
    available = bottom - BODY_TOP - 0.05
    size = 13.0
    while size > 8 and sum(_table_height(_table_content(rows, deck, size), widths)) > available:
        size -= 0.5
    content = _table_content(rows, deck, size)
    heights = _table_height(content, widths)
    frame = slide.shapes.add_table(len(content), len(TABLE_COLS), _emu(MARGIN), _emu(BODY_TOP),
                                   _emu(sum(widths)), _emu(sum(heights)))
    frame.name = "BenefitsTable"
    table = frame.table
    table.first_row = True
    for j, w in enumerate(widths):
        table.columns[j].width = _emu(w)
    for i, (row, h) in enumerate(zip(content, heights)):
        table.rows[i].height = _emu(h)
        for j, paras in enumerate(row):
            cell = table.cell(i, j)
            cell.margin_left = cell.margin_right = _emu(CELL_INSET)
            cell.margin_top = cell.margin_bottom = _emu(CELL_INSET)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (ROW_SHADE if i % 2 == 0 else WHITE)
            frame_ = cell.text_frame
            frame_.text = ""
            _fill_frame_cell(frame_, paras)
    return slide, claims


def _fill_frame_cell(frame, paras: list[Para]) -> None:
    frame.word_wrap = True
    for n, para in enumerate(paras):
        p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
        for run in para.runs:
            r = p.add_run()
            r.text = run.text
            font = r.font
            font.name, font.size, font.bold, font.italic = FONT, Pt(run.size), run.bold, run.italic
            font.color.rgb = run.color
            if run.superscript:
                r._r.get_or_add_rPr().set("baseline", "30000")


def _framing_by_selection_claim(deck: _Deck) -> dict[str, list[Claim]]:
    framing: dict[str, list[Claim]] = {}
    for c in deck.ctx.deck.slides[3].bullets:
        if c.metadata.get("framing_of") and deck.shown(c):
            framing.setdefault(c.metadata["framing_of"], []).append(c)
    return framing


def _with_framing(claims: list[Claim], deck: _Deck, framing: dict[str, list[Claim]], size: float,
                  gap: float) -> list[Para]:
    paras = []
    for n, claim in enumerate(claims):
        paras.append(Para(_claim_runs(claim, deck, size), space_before=gap if n else 0, indent=0.28, bullet="•"))
        for sub in framing.get(claim.metadata.get("selection_claim", ""), []):
            paras.append(Para(_claim_runs(sub, deck, size - 2, italic=True, color=GREY), space_before=1,
                              indent=0.28))
    return paras


def _slide4(prs, deck: _Deck, date: str):
    ctx = deck.ctx
    slide = _base_slide(prs, 4, date)
    s4 = ctx.deck.slides[3]
    framing = _framing_by_selection_claim(deck)
    reasons = [c for c in s4.bullets if c.policy_id and deck.shown(c)][:SLIDE4_MAX_POLICY_BULLETS]
    supporting = [c for c in s4.supporting_benefits if deck.shown(c)][:SLIDE4_MAX_SUPPORTING_BENEFITS]
    limitations = [c for c in s4.key_limitations if deck.shown(c)][:SLIDE4_MAX_KEY_LIMITATIONS]
    subs = [f for c in reasons + limitations for f in framing.get(c.metadata.get("selection_claim", ""), [])]
    claims = reasons + subs + supporting + limitations
    for c in reasons + [f for c in reasons for f in framing.get(c.metadata.get("selection_claim", ""), [])] \
            + supporting + limitations:
        deck.marker(c)
    bottom = _slide_footnotes(slide, claims, deck)
    block = ctx.deck.recommended
    addons = ", ".join(block.required_addons) or "none"
    decided = "the advisor" if block.decided_by.value == "ADVISOR" else "LLM selection, validated and audited"
    _textbox(slide, "RecommendedPolicy", MARGIN, BODY_TOP - 0.1, SLIDE_W - 2 * MARGIN, 0.95, [
        Para([Run(block.policy_name, 26, bold=True, color=NAVY)]),
        Para([Run(f"Variant: {block.variant or 'not specified'}   |   Required add-ons: {addons}   |   "
                  f"Selected by: {decided}", 13, color=GREY)], space_before=2)])
    top = BODY_TOP + 0.95
    left_w, gap_w = 7.6, 0.3
    right_x, right_w = MARGIN + left_w + gap_w, SLIDE_W - 2 * MARGIN - left_w - gap_w
    height = bottom - top

    def left(size):
        return [Para([Run(f"Why {block.policy_name}", size + 1, bold=True, color=ACCENT)])] + \
            _with_framing(reasons, deck, framing, size, size * 0.55)

    def right(size):
        paras = []
        if supporting:
            paras.append(Para([Run("Supporting benefits", size + 1, bold=True, color=ACCENT)]))
            paras += _with_framing(supporting, deck, framing, size, size * 0.4)
        if limitations:
            paras.append(Para([Run("Key limitations", size + 1, bold=True, color=ACCENT)],
                              space_before=size if supporting else 0))
            paras += _with_framing(limitations, deck, framing, size, size * 0.4)
        return paras

    left_paras, _ = _fit(left, left_w, height, 16, 9)
    right_paras, _ = _fit(right, right_w, height, 14, 8.5)
    _textbox(slide, "Reasons", MARGIN, top, left_w, height, left_paras)
    _rect(slide, right_x - gap_w / 2 - 0.01, top + 0.05, 0.02, height - 0.1, ROW_SHADE, "Divider")
    _textbox(slide, "BenefitsAndLimitations", right_x, top, right_w, height, right_paras)
    return slide, claims


def _slide5(prs, deck: _Deck, date: str):
    ctx = deck.ctx
    slide = _base_slide(prs, 5, date)
    s5 = ctx.deck.slides[4]
    assumptions = [c for c in s5.bullets if deck.shown(c)]
    attested = [c for c in ctx.deck.all_claims() if deck.shown(c) and deck.status(c) == AuditStatus.ADVISOR_ATTESTED]
    width = SLIDE_W - 2 * MARGIN
    disclaimer = [Para([Run(ctx.deck.disclaimer, 12, bold=True, italic=True, color=NAVY)])]
    disc_h = paras_height(disclaimer, width)
    _textbox(slide, "Disclaimer", MARGIN, BODY_BOTTOM - disc_h, width, disc_h, disclaimer)
    top, height = BODY_TOP, BODY_BOTTOM - disc_h - 0.1 - BODY_TOP
    left_w, gap_w = 6.6, 0.3
    right_x, right_w = MARGIN + left_w + gap_w, width - left_w - gap_w
    terms = list(dict.fromkeys(to_display(f) for f in s5.footnotes))

    def left(size):
        paras = [Para([Run("Qualifiers and key terms", size + 1, bold=True, color=ACCENT)])]
        paras += [Para([Run(f"[{n}] ", size, bold=True, color=ACCENT), Run(text, size)], space_before=size * 0.3,
                       indent=0.0) for n, text in enumerate(deck.qualifier_texts, start=1)]
        paras += [Para([Run(t, size)], space_before=size * 0.3, indent=0.25, bullet="•") for t in terms]
        if attested:
            paras.append(Para([Run("Advisor-attested", size + 1, bold=True, color=ACCENT)], space_before=size))
            paras += [Para([Run(f"{claim_text(c, deck)} — {deck.results[c.claim_id].advisor_note}", size)],
                           space_before=size * 0.3, indent=0.25, bullet="•") for c in attested]
        return paras

    def right(size):
        paras = [Para([Run("Assumptions", size + 1, bold=True, color=ACCENT)])]
        paras += [Para(_claim_runs(c, deck, size), space_before=size * 0.3, indent=0.25, bullet="•")
                  for c in assumptions]
        paras.append(Para([Run("Sources", size + 1, bold=True, color=ACCENT)], space_before=size))
        paras += [Para([Run(to_display(s), size - 1, color=GREY)], space_before=size * 0.3, indent=0.25,
                       bullet="•") for s in ctx.deck.sources]
        return paras

    left_paras, _ = _fit(left, left_w, height, 13, 7.5)
    right_paras, _ = _fit(right, right_w, height, 13, 7.5)
    _textbox(slide, "QualifiersAndTerms", MARGIN, top, left_w, height, left_paras)
    _textbox(slide, "AssumptionsAndSources", right_x, top, right_w, height, right_paras)
    return slide, list(s5.bullets)


# --- Speaker notes -------------------------------------------------------------------------------------------------


def _evidence_label(evidence_id: str, deck: _Deck) -> str:
    return deck.notes_lookup.get(evidence_id, evidence_id)


def _notes(slide, number: int, deck: _Deck, extra: list[str]) -> None:
    ctx = deck.ctx
    lines = [*extra]
    slide_claims = ctx.deck.slides[number - 1].all_claims()
    facts = {f.fact_id: f for f in ctx.company_profile.facts} if ctx.company_profile else {}
    for claim in slide_claims:
        result = deck.results.get(claim.claim_id)
        status = result.status.value if result else "NOT AUDITED"
        shown = deck.shown(claim)
        refs = [_evidence_label(e, deck) for e in (result.supporting_evidence_ids if result else [])]
        refs += [f"{f} (company profile: {facts[f].display_label})" if f in facts else f
                 for f in (result.supporting_fact_ids if result else [])]
        where = "; ".join(refs) or "no evidence item (code-made or cell-based)"
        lines.append(f"{claim.claim_id} → {status} → {where}" + ("" if shown else " [not rendered]"))
    slide.notes_slide.notes_text_frame.text = "\n".join(lines) or "No claims on this slide."


def _notes_lookup(ctx: RunContext) -> dict[str, str]:
    lookup = {i.evidence_id: f"Marsh profile, section {i.page}" for i in marsh_evidence()}
    policy_ids = [d.document_id for d in ctx.selected_documents]
    try:
        store = load_evidence(policy_ids)
    except (FileNotFoundError, ValueError):
        return lookup
    for p in policy_ids:
        name = store.document(p).display_name
        for item in store.items_for_policy(p, citable_only=False):
            lookup[item.evidence_id] = f"{item.evidence_id} ({name}, p. {item.page})"
    for eid in list(lookup):
        if eid.startswith("EV-MARSH-"):
            lookup[eid] = f"{eid} ({lookup[eid]})"
    return lookup


# --- Entry points -------------------------------------------------------------------------------------------------


def build_pptx(ctx: RunContext, path: Path, *, date: datetime | None = None) -> Path:
    """Draw the 5 slides of the run's audited deck into `path` (no gate; see `render`)."""
    results = {r.claim_id: r for r in ctx.audit_report.results} if ctx.audit_report else {}
    deck = _Deck(ctx=ctx, results=results, notes_lookup=_notes_lookup(ctx))
    prs = Presentation()
    prs.slide_width, prs.slide_height = _emu(SLIDE_W), _emu(SLIDE_H)
    when = f"{(date or datetime.now()):%d %B %Y}".lstrip("0")
    block = ctx.deck.recommended
    extra = {4: [f"Recommended: {block.policy_name} ({block.policy_id}); variant {block.variant or '-'}; add-ons "
                 f"{', '.join(block.required_addons) or 'none'}; decided by {block.decided_by.value}"]}
    for number, draw in enumerate((_slide1, _slide2, _slide3, _slide4, _slide5), start=1):
        slide, _ = draw(prs, deck, when)
        _notes(slide, number, deck, extra.get(number, []))
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def render(ctx: RunContext, *, date: datetime | None = None) -> tuple[Path, GateResult]:
    """Gate, then render outputs/<run_id>/pitch.pptx and QA it. Refuses (RenderRefusedError) if the gate FAILs.
    REVIEW_REQUIRED renders (the advisor reviews the deck); export needs gate.export_allowed."""
    gate = run_gate(ctx)
    if gate.status == OverallFlag.FAIL:
        log_decision(ctx.run_id, "render_refused", {"failures": [f.model_dump() for f in gate.failures]})
        raise RenderRefusedError(gate)
    path = build_pptx(ctx, run_dir(ctx.run_id) / PPTX_FILE, date=date)
    structural_qa(path)
    save_json(ctx.deck, run_dir(ctx.run_id) / PITCH_FILE)  # the gate's WM removals
    save_run_context(ctx)
    log_decision(ctx.run_id, "rendered", {"path": str(path), "gate": gate.status.value,
                                          "export_allowed": gate.export_allowed})
    return path, gate


# --- Structural QA --------------------------------------------------------------------------------------------------


def _shape_paras(frame) -> list[tuple[str, float, bool, bool, float, float]]:
    out = []
    for p in frame.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            out.append(("", 10.0, False, False, 0.0, 0.0))
            continue
        size = max((r.font.size.pt if r.font.size else 18.0) for r in runs)
        p_pr = p._p.pPr
        indent = int(p_pr.get("marL", "0")) / 914400 if p_pr is not None else 0.0
        before = p.space_before.pt if p.space_before is not None else 0.0
        out.append((p.text, size, any(bool(r.font.bold) for r in runs), all(bool(r.font.italic) for r in runs),
                    before, indent))
    return out


def _measure(frame, width_in: float) -> float:
    total = sum(para_height(t, s, b, i, before, indent, width_in) for t, s, b, i, before, indent in _shape_paras(frame))
    return total + (frame.margin_top + frame.margin_bottom) / 914400


def _bullets(frame) -> int:
    return sum(1 for p in frame.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("a:buChar")) is not None)


def structural_qa(path: str | Path) -> None:
    """Reopen the rendered file and check it (see the module docstring). Raises RenderQAError."""
    path = Path(path)
    problems: list[str] = []
    if not path.exists() or path.stat().st_size <= MIN_FILE_BYTES:
        raise RenderQAError([f"{path.name} is missing or not larger than {MIN_FILE_BYTES // 1024} KB"])
    prs = Presentation(path)
    if len(prs.slides) != 5:
        problems.append(f"{len(prs.slides)} slides, expected 5")
    limits = {1: SLIDE1_MAX_BULLETS, 2: SLIDE2_MAX_BULLETS}
    for n, slide in enumerate(prs.slides, start=1):
        shapes = {s.name: s for s in slide.shapes}
        title = shapes.get("Title")
        expected = SLIDE_TITLES[n - 1] if n <= len(SLIDE_TITLES) else "?"
        if title is None or title.text_frame.text != expected:
            problems.append(f"slide {n}: title {title.text_frame.text if title else None!r}, expected {expected!r}")
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text)
                width = (shape.width - shape.text_frame.margin_left - shape.text_frame.margin_right) / 914400
                need = _measure(shape.text_frame, width + 2 * INSET) - 0.0
                have = shape.height / 914400
                if need > have * 1.02 + 0.01:
                    problems.append(f"slide {n}: text box {shape.name!r} overflows ({need:.2f} in of text in "
                                    f"{have:.2f} in)")
                if shape.top / 914400 + have > SLIDE_H + 0.01:
                    problems.append(f"slide {n}: {shape.name!r} runs off the slide")
            if shape.has_table:
                table = shape.table
                widths = [c.width / 914400 for c in table.columns]
                need = 0.0
                for row in table.rows:
                    cells = []
                    for j, cell in enumerate(row.cells):
                        texts.append(cell.text_frame.text)
                        inset = (cell.margin_left + cell.margin_right) / 914400
                        cells.append(sum(para_height(t, s, b, i, before, indent, widths[j] - inset)
                                         for t, s, b, i, before, indent in _shape_paras(cell.text_frame))
                                     + (cell.margin_top + cell.margin_bottom) / 914400)
                    need += max(cells)
                room = BODY_BOTTOM - shape.top / 914400
                if need > room + 0.01:
                    problems.append(f"slide {n}: table overflows ({need:.2f} in of rows in {room:.2f} in)")
                if n == 3 and len(table.rows) - 1 > SLIDE3_MAX_ROWS:
                    problems.append(f"slide 3: {len(table.rows) - 1} rows (max {SLIDE3_MAX_ROWS})")
        joined = "\n".join(texts)
        if "{{" in joined:
            problems.append(f"slide {n}: a template placeholder '{{{{' is left")
        if re.search(r"`\s?\d", joined):
            problems.append(f"slide {n}: a backtick before a digit (write ₹)")
        if re.search(r"unverified|MODEL_KNOWLEDGE", joined, re.IGNORECASE):
            problems.append(f"slide {n}: 'Unverified' / a raw fact status is shown (two labels only: Web-sourced, "
                            f"Assumption)")
        if n in limits and "Body" in shapes and _bullets(shapes["Body"].text_frame) > limits[n]:
            problems.append(f"slide {n}: {_bullets(shapes['Body'].text_frame)} bullets (max {limits[n]})")
        if n == 4:
            block = shapes.get("RecommendedPolicy")
            if block is None:
                problems.append("slide 4: no recommended policy block")
            else:
                name_line = block.text_frame.paragraphs[0].text
                named = named_policies(block.text_frame.text)
                if not name_line.strip() or len(named) > 1:
                    problems.append(f"slide 4: the recommendation block must name exactly one policy "
                                    f"(found {sorted(named) or name_line!r})")
            if "Reasons" in shapes and _bullets(shapes["Reasons"].text_frame) > SLIDE4_MAX_POLICY_BULLETS:
                problems.append(f"slide 4: more than {SLIDE4_MAX_POLICY_BULLETS} visible reason bullets")
            if "BenefitsAndLimitations" in shapes and _bullets(shapes["BenefitsAndLimitations"].text_frame) > \
                    SLIDE4_MAX_SUPPORTING_BENEFITS + SLIDE4_MAX_KEY_LIMITATIONS:
                problems.append("slide 4: too many supporting benefits / key limitations")
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if not notes.strip():
            problems.append(f"slide {n}: no speaker notes")
    if problems:
        raise RenderQAError(problems)


def pptx_text(path: str | Path) -> list[str]:
    """All text per slide (for tests and scripts)."""
    out = []
    for slide in Presentation(path).slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts += [c.text_frame.text for row in shape.table.rows for c in row.cells]
        out.append("\n".join(parts))
    return out

