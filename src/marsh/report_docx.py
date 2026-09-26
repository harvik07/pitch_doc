"""audit_report.docx: the audit report as a Word document (CLAUDE.md section 8; brief 2.1 / 2.2).

Rendered from the SAME data that audit.export_report writes to audit_report.json (the authoritative audit result):
no audit logic here and nothing added that isn't in the JSON. Layout: title, summary (overall flag, confidence
score, counts), gate failures, review items, repairs, then every claim slide by slide with its status, evidence
(document, page, section, quote, source text), company facts, qualifier, explanation and checks.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x00, 0x0F, 0x47)
GREY = RGBColor(0x4A, 0x50, 0x72)
FONT = "Arial"


def _style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = NAVY
        style.font.bold = True


def _field(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"{label}: ")
    run.bold = True
    p.add_run(value or "—")


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def write_report_docx(data: dict, path: str | Path) -> Path:
    """Write the audit report `data` (the audit_report.json content) as a .docx file."""
    path = Path(path)
    doc = Document()
    _style(doc)
    summary = data["summary"]
    results = data["results"]
    doc.add_heading(f"Audit report — {data['run_id']}", level=0)

    doc.add_heading("Summary", level=1)
    score = summary.get("confidence_score")
    _field(doc, "Overall flag", summary["overall_flag"])
    _field(doc, "Confidence score", f"{score:.0%}" if score is not None else "n/a")
    note = doc.add_paragraph("Confidence score = verified or verified-with-qualifier share of the factual claims "
                             "(non-factual lines and labelled assumptions excluded).")
    note.runs[0].font.color.rgb = GREY
    counts = summary.get("counts") or {}
    _field(doc, "Claims by status", ", ".join(f"{k} {v}" for k, v in sorted(counts.items())) or "none")

    if summary.get("gate_failures"):
        doc.add_heading("Gate failures (no export)", level=2)
        _bullets(doc, summary["gate_failures"])
    if summary.get("review_items"):
        doc.add_heading("Review items (advisor acknowledgement needed)", level=2)
        _bullets(doc, summary["review_items"])

    repaired = [r for r in results if r.get("repair_history")]
    if repaired:
        doc.add_heading("Repairs (max 2 attempts per failing claim)", level=2)
        for r in repaired:
            for a in r["repair_history"]:
                after = a.get("text_after") or "no true rewrite"
                doc.add_paragraph(f"{r['claim_id']} attempt {a['attempt']}: {a['status_before']} “{a['text_before']}” → "
                                  f"“{after}” → {a.get('status_after') or 'not re-audited'}", style="List Bullet")

    doc.add_heading("Claims", level=1)
    for slide in sorted({r["claim"]["slide"] for r in results}):
        doc.add_heading(f"Slide {slide}", level=2)
        for r in [x for x in results if x["claim"]["slide"] == slide]:
            claim = r["claim"]
            removed = " (removed)" if claim["state"] == "REMOVED" else ""
            doc.add_heading(f"{r['claim_id']} · {r['status']}{removed}", level=3)
            _field(doc, "Claim", claim.get("display_text") or claim["text"])
            _field(doc, "Type", claim["claim_type"] + (f" · policy {claim['policy_id']}" if claim.get("policy_id") else "")
                   + (" · material" if claim.get("material") else ""))
            for e in r.get("evidence") or []:
                if e.get("missing"):
                    _field(doc, "Evidence", f"{e['evidence_id']} (not found)")
                    continue
                where = f"{e['document']}, p. {e['page']}, {e['section']}"
                if e.get("row_label"):
                    where += f", row “{e['row_label']}”"
                _field(doc, "Evidence", f"{where} ({e['evidence_id']})")
                _field(doc, "Source text", e.get("display_text") or e["text"])
            for f in r.get("facts") or []:
                sources = ", ".join(f.get("sources") or []) or "no web source"
                _field(doc, "Company fact", f"{f['fact_id']} ({f['label']}): {f['value']} — {sources}")
            if r.get("quotes"):
                _field(doc, "Quote", " / ".join(r["quotes"]))
            if r.get("required_qualifier"):
                _field(doc, "Qualifier", r["required_qualifier"])
            _field(doc, "Explanation", r.get("explanation") or "")
            checks = [f"{c['name']}: {c['result']}" + (f" ({c['details']})" if c.get("details") else "")
                      for c in r.get("checks") or []]
            if checks:
                _field(doc, "Checks", "; ".join(checks))
            if r.get("advisor_action"):
                _field(doc, "Advisor", r["advisor_action"] + (f" — {r['advisor_note']}" if r.get("advisor_note") else ""))
    doc.add_paragraph().add_run().add_break(WD_BREAK.LINE)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
