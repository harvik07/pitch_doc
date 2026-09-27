"""audit.py: the independent audit (mocked audit LLM, committed evidence and coverage matrices).

Policy facts used here are CLAUDE.md section 2 golden facts or quoted from the extracted evidence.
"""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone

import pytest

from marsh import api, audit, settings
from marsh.company import build_profile
from marsh.llm import LLMCallError
from marsh.models import (
    FactStatus,
    WebSource,
    AuditResponse,
    AuditStatus,
    Claim,
    ClaimState,
    ClaimType,
    CompanyProfileResponse,
    OverallFlag,
    PitchSlide,
)
from marsh.marsh_profile import load_profile

REAL_CACHE_DIR = settings.CACHE_DIR
NIVA_AIR = "Air Ambulance: up to INR 2,50,000 per Hospitalisation"  # golden fact, EV-NIVA-2-015
NIVA_FOOTNOTE_8 = "Maximum coverage offered for 30 days/policy year/insured person"  # EV-NIVA-2-073


def fact(field, value, status="MODEL_KNOWLEDGE", **web):
    return {"field": field, "value": value, "status": status, "confidence": "medium", "rationale": "Placeholder.",
            **web}


# A placeholder page about the fictional "Example Co" (not a claim about any real company).
WEB_QUOTE = "Example Co employs more than 200,000 people"
WEB_SOURCES = [WebSource(source_id="WEB-001", url="https://example.com/about", title="About Example Co",
                         retrieved_at=datetime(2026, 9, 26, tzinfo=timezone.utc),
                         content=f"About us. {WEB_QUOTE} in placeholder offices.")]
PROFILE = build_profile("Example Co", CompanyProfileResponse.model_validate({"company_recognised": True, "facts": [
    fact("industry", "Placeholder industry services"), fact("size", "Very large enterprise"),
    fact("headcount_band", "200,000+ employees", "WEB_SOURCED", source_ids=["WEB-001"], quotes=[WEB_QUOTE]),
    fact("business_risk", "Client concentration"), fact("business_risk", "Talent attrition"),
    fact("workforce_profile", "Predominantly desk-based professionals", "ASSUMPTION"),
]}), WEB_SOURCES)


@pytest.fixture(scope="module")
def sources():
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(settings, "CACHE_DIR", REAL_CACHE_DIR)
        return audit.load_sources(list(settings.BUNDLED_POLICY_FILES), profile=PROFILE)


class FakeAuditor:
    """The audit LLM: a verdict per claim text (default: UNSUPPORTED, no evidence). Records every call."""

    def __init__(self, verdicts=None, fail=False):
        self.verdicts, self.fail, self.calls = verdicts or {}, fail, []

    def __call__(self, prompt_name, variables, response_model, model=None, run_id=None, **kwargs):
        self.calls.append({"prompt": prompt_name, "variables": variables, "model": model})
        if self.fail:
            raise LLMCallError("timeout")
        out = []
        for line in variables["claims"].splitlines():
            claim_id, _, _, text = [p.strip() for p in line.lstrip("- ").split(" | ", 3)]
            verdict = self.verdicts.get(text, {"status": "UNSUPPORTED", "explanation": "Not in the evidence."})
            if verdict is not None:
                out.append({"claim_id": claim_id, **verdict})
        return AuditResponse.model_validate({"verdicts": out})


def verified(*ids, quote=None, qualifier=None, status="VERIFIED"):
    return {"status": status, "supporting_evidence_ids": list(ids), "explanation": "The evidence states it.",
            "quotes": [{"evidence_id": ids[0], "quote": quote}] if quote else [], "required_qualifier": qualifier}


def claim(text, policy="POL-NIVA", claim_type=ClaimType.POLICY_BENEFIT, n=1, slide=4, **kw):
    return Claim(claim_id=f"CL-{n:03d}", slide_number=slide, text=text, claim_type=claim_type, policy_id=policy, **kw)


def run(monkeypatch, sources, claims, verdicts=None, **fake):
    auditor = FakeAuditor(verdicts, **fake)
    monkeypatch.setattr(audit, "call_structured", auditor)
    results = audit.audit_claims(claims, sources)
    return results, auditor


def check(result, name):
    return next(c for c in result.checks if c.name == name)


# --- Independence and batching --------------------------------------------------------------------------------


def test_the_auditor_never_sees_the_generators_citations(monkeypatch, sources):
    c = claim(f"{NIVA_AIR}.", cited_evidence_ids=["EV-NIVA-2-015"],
              metadata={"selection_claim": "SC-1", "selection_quote": "a generator quote"})
    _, auditor = run(monkeypatch, sources, [c])
    variables = auditor.calls[0]["variables"]
    assert auditor.calls[0]["model"] == settings.GEMINI_AUDIT_MODEL
    assert variables["claims"] == f"- CL-001 | slide 4 | POLICY_BENEFIT | {NIVA_AIR}."
    assert "SC-1" not in json.dumps(variables) and "a generator quote" not in json.dumps(variables)
    assert "EV-NIVA-2-015 | 2 |" in variables["evidence"]  # the full evidence: the auditor finds its own support
    assert "EV-HDFC" not in variables["evidence"]


def test_one_call_per_policy(monkeypatch, sources):
    claims = [claim("A.", n=1), claim("B.", "POL-HDFC", n=2), claim("C.", n=3)]
    results, auditor = run(monkeypatch, sources, claims)
    assert sorted(c["variables"]["policy_id"] for c in auditor.calls) == ["POL-HDFC", "POL-NIVA"]  # concurrent calls
    assert "CL-003" in auditor.calls[0]["variables"]["claims"] and list(results) == ["CL-001", "CL-002", "CL-003"]


def test_a_missing_verdict_is_asked_again_then_needs_review(monkeypatch, sources):
    results, auditor = run(monkeypatch, sources, [claim("Skipped.")], {"Skipped.": None})
    assert len(auditor.calls) == 2 and results["CL-001"].status == AuditStatus.NEEDS_REVIEW
    results, _ = run(monkeypatch, sources, [claim("Any.")], fail=True)
    assert results["CL-001"].status == AuditStatus.NEEDS_REVIEW and "LLM call failed" in results["CL-001"].explanation


# --- 1 ownership, 2 quotes, 3 numbers, 4 policy name ----------------------------------------------------------


def test_ownership(monkeypatch, sources):
    texts = ["Air: Up to INR 5,00,000.", "Road ambulance is covered.", NIVA_AIR]
    verdicts = {texts[0]: verified("EV-HDFC-11-029", quote="Air: Up to INR 5,00,000"),  # another policy's item
                texts[1]: verified("EV-NIVA-9-999"),  # unknown id: dropped, nothing left
                texts[2]: verified("EV-NIVA-2-015", "EV-NIVA-9-999", quote=NIVA_AIR)}
    results, _ = run(monkeypatch, sources, [claim(t, n=n) for n, t in enumerate(texts, 1)], verdicts)
    assert results["CL-001"].status == AuditStatus.UNSUPPORTED and "another policy" in results["CL-001"].explanation
    assert results["CL-002"].status == AuditStatus.UNSUPPORTED
    assert results["CL-003"].status == AuditStatus.VERIFIED
    assert results["CL-003"].supporting_evidence_ids == ["EV-NIVA-2-015"]


def test_quotes_must_be_verbatim(monkeypatch, sources):
    texts = [NIVA_AIR, "No sub-limits on modern or conventional treatments."]
    verdicts = {texts[0]: verified("EV-NIVA-2-015", quote="Air Ambulance: up to INR 2,50,000 per trip"),
                texts[1]: verified("EV-CARE-1-015", quote="No sub-limits on Modern or Conventional Treatments")}
    results, _ = run(monkeypatch, sources, [claim(texts[0]), claim(texts[1], "POL-CARE", n=2)], verdicts)
    assert results["CL-001"].status == AuditStatus.NEEDS_REVIEW and results["CL-001"].quote_check.value == "FAIL"
    assert results["CL-002"].status == AuditStatus.VERIFIED  # heading + continuation rule


def test_number_check_overrides_the_llm(monkeypatch, sources):
    texts = ["Air ambulance is covered up to ₹5,00,000 per hospitalisation.", "Air ambulance is covered for 3 days."]
    verdicts = {t: verified("EV-NIVA-2-015", quote=NIVA_AIR) for t in texts}
    results, _ = run(monkeypatch, sources, [claim(t, n=n) for n, t in enumerate(texts, 1)], verdicts)
    assert results["CL-001"].status == AuditStatus.CONTRADICTED and results["CL-001"].llm_status == AuditStatus.VERIFIED
    assert results["CL-001"].number_check.result.value == "FAIL"
    assert results["CL-002"].status == AuditStatus.UNSUPPORTED  # no DAYS number in the supporting item


def test_label_numbers_count_only_for_the_claims_topic(monkeypatch, sources):
    texts = ["HDFC ERGO Optima Secure+ offers a ₹25 lakh deductible.",
             "Shared accommodation pays ₹800 per day, up to ₹4,800, for base sum insured up to ₹15 lakh."]
    verdicts = {texts[0]: verified("EV-HDFC-7-006", quote="15%"),  # column "Base SI = 25 Lakhs"
                texts[1]: verified("EV-NIVA-2-031", quote="INR 800 per day; Maximum INR 4,800")}
    results, _ = run(monkeypatch, sources, [claim(texts[0], "POL-HDFC"), claim(texts[1], n=2)], verdicts)
    assert results["CL-001"].status in (AuditStatus.UNSUPPORTED, AuditStatus.CONTRADICTED)
    assert results["CL-002"].status == AuditStatus.VERIFIED  # golden fact, SI tier stated in the claim


def test_policy_name_check(monkeypatch, sources):
    text = "HDFC ERGO Optima Secure+ covers air ambulance up to ₹2,50,000 per hospitalisation."
    results, _ = run(monkeypatch, sources, [claim(text)], {text: verified("EV-NIVA-2-015", quote=NIVA_AIR)})
    r = results["CL-001"]
    assert r.status == AuditStatus.UNSUPPORTED and check(r, "policy_name").result.value == "FAIL"
    summary = audit.build_summary([r], [claim(text)])
    assert any("wrong policy reference" in f for f in summary.gate_failures)


# --- 5 topic anchor, 6 absolute language ----------------------------------------------------------------------


def test_topic_anchor(monkeypatch, sources):
    texts = ["Niva Bupa ReAssure 2.0 has a 30-day initial waiting period.",
             "Hospital cash is paid for up to 30 days per policy year."]
    verdicts = {t: verified("EV-NIVA-2-073", quote=NIVA_FOOTNOTE_8) for t in texts}
    results, _ = run(monkeypatch, sources, [claim(t, n=n) for n, t in enumerate(texts, 1)], verdicts)
    assert results["CL-001"].status == AuditStatus.UNSUPPORTED  # the 30 days are a hospital-cash limit
    assert check(results["CL-001"], "topic_anchor").result.value == "FAIL"
    assert results["CL-002"].status == AuditStatus.VERIFIED  # a footnote's topic includes the rows that cite it


def test_absolute_language(monkeypatch, sources):
    texts = ["Aditya Birla Health Activ One guarantees 100% HealthReturns every year.",
             "Unlimited e-consultation is available within the network."]
    verdicts = {texts[0]: verified("EV-ABHI-1-012", quote="100% HealthReturns TM*"),
                texts[1]: verified("EV-NIVA-2-036", quote="Unlimited e-consultation within our network.")}
    results, _ = run(monkeypatch, sources, [claim(texts[0], "POL-ABHI"), claim(texts[1], n=2)], verdicts)
    assert results["CL-001"].status == AuditStatus.NEEDS_REVIEW and "indicative" in results["CL-001"].explanation
    assert results["CL-002"].status == AuditStatus.VERIFIED  # the evidence says "Unlimited"


# --- 7 qualifiers ---------------------------------------------------------------------------------------------


def test_structured_and_footnote_qualifiers(monkeypatch, sources):
    texts = ["Unutilised base sum insured carries forward, up to 10 times of base sum insured.",
             "Titanium+ carries unutilised base sum insured forward, up to 10 times of base sum insured.",
             "Air ambulance is covered up to ₹5 lacs per year.",
             "Hospitalisation is covered for 2 hours and more."]
    quote = "10X: Unutilised Base Sum Insured carries forward to the next policy year"
    verdicts = {texts[0]: verified("EV-NIVA-2-027", quote=quote), texts[1]: verified("EV-NIVA-2-027", quote=quote),
                texts[2]: verified("EV-CARE-3-012", quote="Up to `5 lacs per year"),
                texts[3]: verified("EV-NIVA-1-041", quote="Hospitalisation covered for 2 hours and more")}
    policies = ["POL-NIVA", "POL-NIVA", "POL-CARE", "POL-NIVA"]
    results, _ = run(monkeypatch, sources, [claim(t, p, n=n) for n, (t, p) in enumerate(zip(texts, policies), 1)],
                     verdicts)
    assert results["CL-001"].status == AuditStatus.VERIFIED_WITH_QUALIFIER
    assert "Applies to Titanium+ only" in results["CL-001"].required_qualifier
    assert results["CL-002"].status == AuditStatus.VERIFIED  # the claim names its variant
    assert "Optional benefit at extra premium" in results["CL-003"].required_qualifier  # CARE air ambulance
    assert "24 hours" in results["CL-004"].required_qualifier  # footnote (11), AYUSH


def test_a_qualifier_with_a_wrong_number_needs_review(monkeypatch, sources):
    verdicts = {NIVA_AIR: verified("EV-NIVA-2-015", quote=NIVA_AIR, qualifier="Only for the first 45 days")}
    results, _ = run(monkeypatch, sources, [claim(NIVA_AIR)], verdicts)
    assert results["CL-001"].status == AuditStatus.NEEDS_REVIEW


# --- 8 absence claims -----------------------------------------------------------------------------------------


def test_absence_claims(monkeypatch, sources):
    texts = ["Treatment abroad is not stated in the Niva Bupa ReAssure 2.0 brochure.",
             "Air ambulance is not stated in the Niva Bupa ReAssure 2.0 brochure.",
             "Care Supreme does not cover international treatment.",
             "Niva Bupa ReAssure 2.0 does not cover air ambulance.",
             "Maternity is excluded from the HDFC ERGO Optima Secure+ base plan."]
    policies = ["POL-NIVA", "POL-NIVA", "POL-CARE", "POL-NIVA", "POL-HDFC"]
    verdicts = {texts[4]: verified("EV-HDFC-14-011", quote="Maternity")}
    claims = [claim(t, p, n=n) for n, (t, p) in enumerate(zip(texts, policies), 1)]
    results, _ = run(monkeypatch, sources, claims, verdicts)
    assert results["CL-001"].status == AuditStatus.VERIFIED  # the NOT_STATED cell is its evidence
    assert results["CL-002"].status == AuditStatus.CONTRADICTED  # the brochure states air ambulance
    assert results["CL-003"].status == AuditStatus.UNSUPPORTED  # NOT_STATED ≠ excluded
    assert "not stated in the Care Health Care Supreme brochure" in results["CL-003"].explanation
    assert results["CL-004"].status == AuditStatus.CONTRADICTED
    assert check(results["CL-005"], "absence").result.value == "PASS"
    assert "add-on" in (results["CL-005"].required_qualifier or "")  # the Parenthood add-on covers it


def test_not_stated_with_contrary_evidence_needs_review(monkeypatch, sources):
    text = "Treatment abroad is not stated in the Niva Bupa ReAssure 2.0 brochure."
    verdicts = {text: verified("EV-NIVA-2-015", quote=NIVA_AIR, status="CONTRADICTED")}
    results, _ = run(monkeypatch, sources, [claim(text)], verdicts)
    assert results["CL-001"].status == AuditStatus.NEEDS_REVIEW


def test_code_not_stated_rows_need_no_llm(monkeypatch, sources):
    row = claim("Not stated in the brochure", claim_type=ClaimType.POLICY_FACT, slide=3,
                metadata={"match_id": "MATCH-NIVA-MATERNITY"})
    results, auditor = run(monkeypatch, sources, [row])
    assert results["CL-001"].status == AuditStatus.VERIFIED and auditor.calls == []


# --- 9 slide 2, 10 company claims, 11 pricing, NON_FACTUAL -----------------------------------------------------


def marsh_verdict(ms_id, quote=None):
    """The audit LLM's verdict for a Marsh statement: the record's evidence item, quoting its fact."""
    record = load_profile().record(ms_id)
    item = next(i for i in audit.marsh_evidence() if i.row_label == ms_id)
    return verified(item.evidence_id, quote=quote or record.fact)


def test_slide2(monkeypatch, sources):
    profile = load_profile()
    ms021, ms026 = profile.record("MS-021"), profile.record("MS-026")
    without = ms026.fact.replace("generally ", "")
    invented = "Marsh is the world's leading broker."
    claims = [claim(ms021.fact, None, ClaimType.MARSH_STATEMENT, n=1, slide=2),
              claim(invented, None, ClaimType.MARSH_STATEMENT, n=2, slide=2),
              claim("98% health claims payout ratio.", "POL-HDFC", ClaimType.POLICY_FACT, n=3, slide=2),
              claim("A risk partner for your workforce", None, ClaimType.NON_FACTUAL, n=4, slide=2, material=False),
              claim(without, None, ClaimType.MARSH_STATEMENT, n=5, slide=2),
              claim(ms021.fact, None, ClaimType.MARSH_STATEMENT, n=6, slide=2),
              claim("Its placeholder industry makes this relevant.", None, ClaimType.NON_FACTUAL, n=7, slide=2,
                    material=False, basis_fact_ids=["CF-001"]),
              claim("Its 500 offices make this relevant.", None, ClaimType.NON_FACTUAL, n=8, slide=2,
                    material=False, basis_fact_ids=["CF-001"])]
    verdicts = {ms021.fact: marsh_verdict("MS-021"), without: marsh_verdict("MS-026", quote="as an independent "
                                                                                           "insurance intermediary"),
                invented: verified("EV-NIVA-2-015", quote=NIVA_AIR)}
    results, auditor = run(monkeypatch, sources, claims, verdicts)
    assert [c["variables"]["policy_id"] for c in auditor.calls] == ["MARSH"]  # one batch against the profile
    assert "EV-NIVA" not in auditor.calls[0]["variables"]["evidence"]
    assert results["CL-001"].status == AuditStatus.VERIFIED
    assert results["CL-001"].supporting_evidence_ids[0].startswith("EV-MARSH-2-")
    assert results["CL-002"].status == AuditStatus.UNSUPPORTED  # a policy's evidence can't support Marsh
    assert results["CL-003"].status == AuditStatus.UNSUPPORTED  # an insurer statistic on slide 2
    assert results["CL-004"].status == AuditStatus.NON_FACTUAL  # the headline
    assert results["CL-005"].status == AuditStatus.NEEDS_REVIEW  # drops "generally" (its source's condition)
    assert check(results["CL-005"], "condition").result.value == "FAIL"
    assert results["CL-007"].status == AuditStatus.NON_FACTUAL  # why it matters: rests on a company fact
    assert results["CL-008"].status == AuditStatus.NEEDS_REVIEW  # a number that isn't in its company fact
    summary = audit.build_summary(list(results.values()), claims)
    assert any("non-Marsh claim on slide 2" in f for f in summary.gate_failures)


def test_a_marsh_statement_needs_a_verbatim_quote(monkeypatch, sources):
    fact = load_profile().record("MS-021").fact
    results, _ = run(monkeypatch, sources, [claim(fact, None, ClaimType.MARSH_STATEMENT, slide=2)],
                     {fact: marsh_verdict("MS-021", quote="Marsh invents a quote here")})
    assert results["CL-001"].status == AuditStatus.NEEDS_REVIEW


def company(text, basis, n=1, slide=1, claim_type=ClaimType.COMPANY_FACT, label="web"):
    return Claim(claim_id=f"CL-{n:03d}", slide_number=slide, text=text, claim_type=claim_type,
                 basis_fact_ids=basis, material=False, qualifier_text="Web-sourced" if label == "web" else None)


def test_company_claims_are_deterministic(monkeypatch, sources):
    assert PROFILE.facts[2].status == FactStatus.WEB_SOURCED and PROFILE.facts[1].status == FactStatus.MODEL_KNOWLEDGE
    # Company Overview statements are worded as their facts (CF-003 "200,000+ employees", CF-002 "Very large
    # enterprise", CF-006 "Predominantly desk-based professionals")
    claims = [company("Over 200,000 employees.", ["CF-003"], 1),
              company("214,356 employees.", ["CF-003"], 2),
              company("Over 300,000 employees.", ["CF-003"], 3),
              company("A very large enterprise.", ["CF-099"], 4),
              company("A very large enterprise.", ["CF-002"], 5, label=None),
              company("Predominantly desk-based professionals.*", ["CF-006"], 6, claim_type=ClaimType.ASSUMPTION,
                      label=None),
              company("Predominantly desk-based professionals.", ["CF-006"], 7, claim_type=ClaimType.ASSUMPTION,
                      label=None),
              company("Client concentration is a key business risk.*", ["CF-004"], 8, slide=4,
                      claim_type=ClaimType.ASSUMPTION, label=None),
              company("A very large enterprise.*", ["CF-002"], 9, claim_type=ClaimType.ASSUMPTION,
                      label=None),
              company("A very large enterprise.", ["CF-002"], 10),  # "Web-sourced" on a model-knowledge fact
              company("The company employs over 200,000 people in Bengaluru.", ["CF-003"], 11)]  # adds wording
    results, auditor = run(monkeypatch, sources, claims)
    status = {k: r.status for k, r in results.items()}
    assert auditor.calls == []
    assert status["CL-001"] == AuditStatus.VERIFIED and results["CL-001"].supporting_fact_ids == ["CF-003"]
    assert check(results["CL-001"], "web_source").result.value == "PASS"
    assert check(results["CL-001"], "wording").result.value == "PASS"
    assert status["CL-002"] == AuditStatus.CONTRADICTED  # 214,356 isn't its fact's number (200,000+)
    assert status["CL-003"] == AuditStatus.CONTRADICTED  # 300,000 vs the profile's 200,000+
    assert status["CL-004"] == AuditStatus.UNSUPPORTED
    assert status["CL-005"] == AuditStatus.NEEDS_REVIEW  # neither label
    assert status["CL-006"] == AuditStatus.LABELLED_ASSUMPTION
    assert status["CL-007"] == AuditStatus.NEEDS_REVIEW  # no "*" assumption marker
    assert status["CL-008"] == AuditStatus.NEEDS_REVIEW  # business risk outside slide 1
    assert status["CL-009"] == AuditStatus.LABELLED_ASSUMPTION  # MODEL_KNOWLEDGE is shown as an assumption
    assert status["CL-010"] == AuditStatus.NEEDS_REVIEW and "not web-sourced" in results["CL-010"].explanation
    # a web-sourced fact can't carry words (a place, "people") its fact doesn't state onto the Company Overview
    assert status["CL-011"] == AuditStatus.NEEDS_REVIEW and check(results["CL-011"], "wording").result.value == "FAIL"
    assert "Bengaluru" in results["CL-011"].explanation
    no_profile = audit.AuditSources(store=sources.store, cells=sources.cells, taxonomy=sources.taxonomy,
                                    topics=sources.topics)
    assert audit.audit_claims([claims[0]], no_profile)["CL-001"].status == AuditStatus.NEEDS_REVIEW


def test_web_sourced_company_claims_are_re_checked_against_the_pages(monkeypatch, sources):
    """The audit doesn't trust the stored WEB_SOURCED status: a quote no longer in the page → NEEDS_REVIEW."""
    tampered = PROFILE.model_copy(deep=True)
    tampered.sources[0].content = "About us. A different page."
    changed = audit.AuditSources(store=sources.store, cells=sources.cells, profile=tampered,
                                 taxonomy=sources.taxonomy, topics=sources.topics)
    monkeypatch.setattr(audit, "call_structured", FakeAuditor())
    result = audit.audit_claims([company("The company employs over 200,000 people.", ["CF-003"])], changed)["CL-001"]
    assert result.status == AuditStatus.NEEDS_REVIEW and "quote not found" in result.explanation


def test_the_audit_report_shows_fact_labels_never_raw_statuses(monkeypatch, sources):
    claims = [company("The company employs over 200,000 people.", ["CF-003"], 1),
              company("It is a very large enterprise.*", ["CF-002"], 2, claim_type=ClaimType.ASSUMPTION,
                      label=None)]
    slide = PitchSlide(slide_number=1, title="Company Overview", bullets=claims)
    run_sources = audit.AuditSources(store=sources.store, cells=sources.cells, profile=PROFILE,
                                     run_id="RUN-20260926-000000-0002", taxonomy=sources.taxonomy,
                                     topics=sources.topics)
    monkeypatch.setattr(audit, "call_structured", FakeAuditor())
    report = audit.audit_deck([slide], run_sources, "RUN-20260926-000000-0002")
    json_path, md_path = audit.export_report(report, claims, run_sources)
    data = json.loads(json_path.read_text(encoding="utf-8"))
    facts = {r["claim_id"]: r["facts"] for r in data["results"]}
    assert facts["CL-001"] == [{"fact_id": "CF-003", "value": "200,000+ employees", "label": "Web-sourced",
                                "sources": ["https://example.com/about"], "quotes": [WEB_QUOTE]}]
    assert facts["CL-002"][0]["label"] == "Assumption" and "status" not in facts["CL-002"][0]
    md = md_path.read_text(encoding="utf-8")
    assert "CF-003 (Web-sourced: https://example.com/about)" in md and "CF-002 (Assumption)" in md
    for shown in (json_path.read_text(encoding="utf-8"), md):
        assert "MODEL_KNOWLEDGE" not in shown and "unverified" not in shown.lower()


def test_precise_figures():
    assert audit.precise_figures("A workforce of 214,356 employees.") == ["214,356"]
    assert audit.precise_figures("Over 200,000 employees and 200,000+ staff.") == []
    assert audit.precise_figures("Revenue of $18.2 billion.") == ["$18.2 billion"]
    assert audit.precise_figures("Founded in 1981 with 7 offices.") == []


def test_pricing(monkeypatch, sources):
    texts = ["The premium is lower than Niva Bupa ReAssure 2.0's premium of ₹22,616.",
             "The premium is ₹22,616.",
             "The premium is ₹22,616 in the brochure's example.",
             "The premium is ₹22,616 for a 2-member family floater aged 35 and 30 years with ₹10 lakh base cover."]
    quote = "INR 22,616 premium is for Optima Secure + plan"  # golden fact, EV-HDFC-3-017 / 3-002
    verdicts = {t: verified("EV-HDFC-3-017", "EV-HDFC-3-002", quote=quote) for t in texts}
    verdicts[texts[2]] = verified("EV-HDFC-3-017", "EV-HDFC-3-002", quote=quote,
                                  qualifier="2-member family floater, age 35 years and 30 years, INR 10 lakhs base cover")
    claims = [claim(t, "POL-HDFC", ClaimType.POLICY_PRICING, n=n) for n, t in enumerate(texts, 1)]
    results, _ = run(monkeypatch, sources, claims, verdicts)
    assert results["CL-001"].status == AuditStatus.UNSUPPORTED and check(results["CL-001"], "pricing").result.value == "FAIL"
    assert results["CL-002"].status == AuditStatus.NEEDS_REVIEW  # no context
    assert results["CL-003"].status == AuditStatus.VERIFIED_WITH_QUALIFIER
    assert check(results["CL-004"], "pricing").result.value == "PASS"


def test_non_factual(monkeypatch, sources):
    claims = [claim("Company details are placeholders.", None, ClaimType.NON_FACTUAL, slide=4, material=False),
              claim(f"{NIVA_AIR} in Niva Bupa ReAssure 2.0.", None, ClaimType.NON_FACTUAL, n=2, slide=4)]
    results, auditor = run(monkeypatch, sources, claims,
                           {f"{NIVA_AIR} in Niva Bupa ReAssure 2.0.": verified("EV-NIVA-2-015", quote=NIVA_AIR)})
    assert results["CL-001"].status == AuditStatus.NON_FACTUAL
    assert results["CL-002"].status == AuditStatus.VERIFIED and len(auditor.calls) == 1  # reclassified, audited


def test_run_assumptions(monkeypatch, sources):
    claims = [claim("Assumed base sum insured: ₹10,00,000*", None, ClaimType.ASSUMPTION),
              claim("Assumed base sum insured: ₹10,00,000", None, ClaimType.ASSUMPTION, n=2),
              claim("Assumed base sum insured: ₹10,00,000 (Assumption)", None, ClaimType.ASSUMPTION, n=3)]
    results, _ = run(monkeypatch, sources, claims)
    assert results["CL-001"].status == AuditStatus.LABELLED_ASSUMPTION
    assert results["CL-002"].status == AuditStatus.NEEDS_REVIEW
    assert results["CL-003"].status == AuditStatus.NEEDS_REVIEW  # the old text label is no longer the marker


# --- Summary, deck, standalone API ----------------------------------------------------------------------------


def test_summary_flags_and_score(monkeypatch, sources):
    texts = [NIVA_AIR, "Hospitalisation is covered for 2 hours and more.", "Home care is covered."]
    verdicts = {texts[0]: verified("EV-NIVA-2-015", quote=NIVA_AIR),
                texts[1]: verified("EV-NIVA-1-041", quote="Hospitalisation covered for 2 hours and more")}
    slide = PitchSlide(slide_number=4, title="Recommended Policy",
                       bullets=[claim(t, n=n) for n, t in enumerate(texts[:2], 1)])
    monkeypatch.setattr(audit, "call_structured", FakeAuditor(verdicts))
    report = audit.audit_deck([slide], sources, "RUN-20260926-000000-0001")
    vwq = slide.bullets[1]
    assert vwq.qualifier_text and vwq.state == ClaimState.AUDITED  # stored for the renderer
    assert report.summary.overall_flag == OverallFlag.PASS and report.summary.confidence_score == 1.0

    slide.bullets[1].qualifier_text = None  # a VWQ whose qualifier isn't rendered
    summary = audit.build_summary(report.results, slide.bullets)
    assert summary.overall_flag == OverallFlag.REVIEW_REQUIRED

    bad = claim(texts[2], n=3)
    results, _ = run(monkeypatch, sources, [bad], verdicts)
    summary = audit.build_summary(report.results + [results["CL-003"]], slide.bullets + [bad])
    assert summary.overall_flag == OverallFlag.FAIL and summary.confidence_score == 0.67
    bad.state = ClaimState.REMOVED
    assert audit.build_summary(report.results + [results["CL-003"]], slide.bullets + [bad]).confidence_score == 1.0


def test_audit_pitch_content_works_standalone(monkeypatch):
    monkeypatch.setattr(settings, "CACHE_DIR", REAL_CACHE_DIR)
    monkeypatch.setattr(audit, "call_structured", FakeAuditor({NIVA_AIR: verified("EV-NIVA-2-015", quote=NIVA_AIR)}))
    slides = [{"slide_number": 4, "title": "Recommended Policy",
               "bullets": [{"claim_id": "CL-001", "slide_number": 4, "text": NIVA_AIR, "claim_type": "POLICY_BENEFIT",
                            "policy_id": "POL-NIVA"}]}]
    report = api.auditPitchContent(slides, ["POL-NIVA"])
    assert report.results[0].status == AuditStatus.VERIFIED and report.summary.overall_flag == OverallFlag.PASS
    out = settings.OUTPUTS_DIR / report.run_id
    data = json.loads((out / "audit_report.json").read_text(encoding="utf-8"))
    row = data["results"][0]
    assert row["claim"]["text"] == NIVA_AIR and row["claim"]["display_text"] == NIVA_AIR
    assert row["evidence"][0]["page"] == 2 and "display_text" in row["evidence"][0]
    md = (out / "audit_report.md").read_text(encoding="utf-8")
    assert md.startswith(f"# Audit report — {report.run_id}") and "**Overall flag: PASS**" in md
    assert re.search(r"\| CL-001 \| 4 \| .* \| VERIFIED \|", md)


def test_a_table_cell_is_audited_in_its_row(monkeypatch, sources):
    cell = claim("Coverage is available only within the insurer network.", claim_type=ClaimType.POLICY_LIMIT,
                 slide=3, metadata={"match_id": "MATCH-NIVA-OPD"})
    _, auditor = run(monkeypatch, sources, [cell])
    assert "| slide 3, table row 'Outpatient consultations' | POLICY_LIMIT |" in auditor.calls[0]["variables"]["claims"]


def test_a_qualifier_stated_in_the_rows_other_cell_is_not_missing(monkeypatch, sources):
    from marsh.models import BenefitRow

    benefit = claim("Hospitalisation is covered for 2 hours and more.", slide=3, metadata={"match_id": "MATCH-NIVA-HOSP"})
    condition = claim("A minimum of 24 hours of hospitalisation is required for AYUSH treatment in an AYUSH hospital.",
                      claim_type=ClaimType.POLICY_LIMIT, n=2, slide=3, metadata={"match_id": "MATCH-NIVA-HOSP"})
    slide = PitchSlide(slide_number=3, title="Policy Benefits Mapped to Exposures", table_rows=[BenefitRow(
        exposure_id="EXP-HOSP", exposure_name="In-patient hospitalisation", benefit=benefit, condition=condition)])
    verdicts = {benefit.text: verified("EV-NIVA-1-041", quote="Hospitalisation covered for 2 hours and more")}
    auditor = FakeAuditor(verdicts)
    monkeypatch.setattr(audit, "call_structured", auditor)
    report = audit.audit_deck([slide], sources, "RUN-20260926-000000-0003")
    assert report.results[0].status == AuditStatus.VERIFIED  # footnote (11) is stated by the condition cell
    assert "the row's other cell: \"A minimum of 24 hours" in auditor.calls[0]["variables"]["claims"]


def test_footnote_sentences_are_not_cut_at_enumerators():
    text = "(6) Eligible person - a. All members except son/daughter & b. Any member aged 18 years. Next sentence."
    assert audit._first_sentence(text) == "Eligible person - a. All members except son/daughter & b. Any member aged 18 years."


def test_re_auditing_an_unchanged_deck_uses_the_cache(monkeypatch, sources):
    import dataclasses

    run_id = "RUN-20260926-000000-0004"
    cached = dataclasses.replace(sources, run_id=run_id, _cache=None)
    texts = [NIVA_AIR, "Home care is covered."]

    def slides():
        return [PitchSlide(slide_number=4, title="Recommended Policy",
                           bullets=[claim(t, n=n) for n, t in enumerate(texts, 1)])]

    auditor = FakeAuditor({NIVA_AIR: verified("EV-NIVA-2-015", quote=NIVA_AIR)})
    monkeypatch.setattr(audit, "call_structured", auditor)
    first = audit.audit_deck(slides(), cached, run_id)
    assert len(auditor.calls) == 1 and (settings.OUTPUTS_DIR / run_id / "audit_cache.json").exists()
    fresh = dataclasses.replace(sources, run_id=run_id, _cache=None)  # a new process: the cache comes from disk
    second = audit.audit_deck(slides(), fresh, run_id)
    assert len(auditor.calls) == 1 and fresh.cache_hits == 2 and second == first  # 0 LLM calls, identical report

    changed = slides()
    changed[0].bullets[1].text = "Home care is covered up to the sum insured."
    audit.audit_deck(changed, fresh, run_id)
    assert len(auditor.calls) == 2 and "CL-001" not in auditor.calls[1]["variables"]["claims"]  # only the new text
    dirty = slides()
    dirty[0].bullets[0].state = ClaimState.DIRTY
    audit.audit_deck(dirty, fresh, run_id)
    assert "CL-001" in auditor.calls[2]["variables"]["claims"]  # a DIRTY claim is always re-audited


def test_the_docx_report_is_the_same_audit_as_the_json(monkeypatch, sources):
    """audit_report.docx is rendered from the audit_report.json data: same claims, statuses and evidence."""
    from docx import Document

    claims = [company("The company employs over 200,000 people.", ["CF-003"], 1, slide=4),
              company("It is a very large enterprise.*", ["CF-002"], 2, slide=4, claim_type=ClaimType.ASSUMPTION,
                      label=None),
              claim(f"{NIVA_AIR}.", n=3)]
    slide = PitchSlide(slide_number=4, title="Recommended Policy", bullets=claims)
    run_sources = audit.AuditSources(store=sources.store, cells=sources.cells, profile=PROFILE,
                                     run_id="RUN-20260926-000000-0003", taxonomy=sources.taxonomy,
                                     topics=sources.topics)
    monkeypatch.setattr(audit, "call_structured", FakeAuditor({f"{NIVA_AIR}.": verified("EV-NIVA-2-015", quote=NIVA_AIR)}))
    report = audit.audit_deck([slide], run_sources, "RUN-20260926-000000-0003")
    json_path, _ = audit.export_report(report, claims, run_sources)
    docx_path = json_path.parent / audit.REPORT_DOCX
    assert docx_path.exists()
    text = "\n".join(p.text for p in Document(docx_path).paragraphs)
    data = json.loads(json_path.read_text(encoding="utf-8"))
    assert f"Overall flag: {data['summary']['overall_flag']}" in text
    for row in data["results"]:
        assert f"{row['claim_id']} · {row['status']}" in text
        assert row["claim"]["display_text"] in text
        for e in row["evidence"]:
            assert e["evidence_id"] in text and e["display_text"] in text
    assert "CF-003 (Web-sourced)" in text and "MODEL_KNOWLEDGE" not in text


# --- Deterministic NON_FACTUAL; grouped, bounded, order-preserving audit calls ------------------------------------


def test_non_factual_lines_need_no_audit_call(monkeypatch, sources):
    lines = [claim("A health and benefits partner for your people", None, ClaimType.NON_FACTUAL, n=1, slide=2,
                   material=False),
             claim("We look forward to working with you.", None, ClaimType.NON_FACTUAL, n=2, material=False)]
    results, auditor = run(monkeypatch, sources, lines)
    assert auditor.calls == [] and {r.status for r in results.values()} == {AuditStatus.NON_FACTUAL}
    numbered = claim("Cover up to ₹5,00,000 is included.", None, ClaimType.NON_FACTUAL, n=3, material=False)
    results, auditor = run(monkeypatch, sources, [numbered])  # a number: not treated as non-factual
    assert results["CL-003"].status != AuditStatus.NON_FACTUAL


def test_audit_calls_are_grouped_share_the_full_evidence_and_are_bounded(monkeypatch, sources):
    import threading
    import time

    monkeypatch.setattr(settings, "AUDIT_CLAIMS_PER_CALL", 2)
    monkeypatch.setattr(settings, "AUDIT_MAX_CONCURRENCY", 2)
    texts = [f"{NIVA_AIR}."] * 5  # five claims (ids differ) with the same true text
    lock, active, peak = threading.Lock(), [0], [0]

    class SlowAuditor(FakeAuditor):
        def __call__(self, *args, **kwargs):
            with lock:
                active[0] += 1
                peak[0] = max(peak[0], active[0])
            time.sleep(0.05)
            try:
                return super().__call__(*args, **kwargs)
            finally:
                with lock:
                    active[0] -= 1

    auditor = SlowAuditor({t: verified("EV-NIVA-2-015", quote=NIVA_AIR) for t in texts})
    monkeypatch.setattr(audit, "call_structured", auditor)
    claims = [claim(t, n=n) for n, t in enumerate(texts, 1)]
    results = audit.audit_claims(claims, sources)
    assert len(auditor.calls) == 3 and 1 < peak[0] <= 2  # 5 claims in groups of 2, 2 at a time
    evidence = {c["variables"]["evidence"] for c in auditor.calls}
    assert len(evidence) == 1  # every group gets the same full evidence set of the policy (section 8)
    full = sources.store.items_for_policy("POL-NIVA")
    assert all(i.evidence_id in next(iter(evidence)) for i in full)
    assert list(results) == [c.claim_id for c in claims]  # claim order, whatever order the calls finished in
    assert {r.status for r in results.values()} == {AuditStatus.VERIFIED}


def test_an_exact_web_sourced_headcount_passes_the_audit(monkeypatch, sources):
    """An exact employee count its checked source states is web-sourced and verified; on an assumption it isn't."""
    page = WebSource(source_id="WEB-001", url="https://exampleco.com/about", title="About Example Co",
                     retrieved_at=datetime(2026, 9, 26, tzinfo=timezone.utc),
                     content="Example Co has 593,798 professionals across placeholder offices.")
    profile = build_profile("Example Co", CompanyProfileResponse.model_validate({"company_recognised": True, "facts": [
        fact("headcount_band", "593,798 professionals", "WEB_SOURCED", source_ids=["WEB-001"],
             quotes=["Example Co has 593,798 professionals"]),
        fact("workforce_profile", "Mostly desk-based staff", "ASSUMPTION")]}), [page])
    assert profile.facts[0].status == FactStatus.WEB_SOURCED
    run_sources = audit.AuditSources(store=sources.store, cells=sources.cells, profile=profile,
                                     taxonomy=sources.taxonomy, topics=sources.topics)
    monkeypatch.setattr(audit, "call_structured", FakeAuditor())
    claims = [company("593,798 professionals.", ["CF-001"], 1),
              company("593,798 professionals.*", ["CF-002"], 2, claim_type=ClaimType.ASSUMPTION, label=None)]
    results = audit.audit_claims(claims, run_sources)
    assert results["CL-001"].status == AuditStatus.VERIFIED
    assert results["CL-002"].status != AuditStatus.LABELLED_ASSUMPTION  # an exact figure isn't an assumption's

