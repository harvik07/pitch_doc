"""render_ppt.py: the fixed template, what is rendered, speaker notes, and structural QA."""

from __future__ import annotations

from datetime import datetime

import pytest
from deck_builder import AIR, SHARED_QUALIFIER, approved_run, set_result
from pptx import Presentation
from pptx.util import Emu

from marsh import settings
from marsh.models import (
    SLIDE_TITLES,
    AdvisorAction,
    AuditStatus,
    ClaimState,
    ClaimType,
    OverallFlag,
    WebSource,
)
from marsh.render_ppt import (
    FONT,
    LEGEND_MANY,
    LEGEND_ONE,
    RenderQAError,
    RenderRefusedError,
    pptx_text,
    render,
    structural_qa,
    text_width_pt,
)

REAL_CACHE_DIR = settings.CACHE_DIR


@pytest.fixture(autouse=True)
def real_evidence(monkeypatch):
    monkeypatch.setattr(settings, "CACHE_DIR", REAL_CACHE_DIR)  # evidence pages for the speaker notes


@pytest.fixture
def rendered():
    ctx = approved_run()
    path, gate = render(ctx)
    return ctx, path, gate


def test_an_approved_deck_renders_and_passes_qa(rendered):
    ctx, path, gate = rendered
    assert gate.status == OverallFlag.PASS and path.name == "pitch.pptx" and path.stat().st_size > 10 * 1024
    assert path.parent == settings.OUTPUTS_DIR / ctx.run_id
    prs = Presentation(path)
    assert len(prs.slides) == 4 and round(prs.slide_width / prs.slide_height, 2) == round(16 / 9, 2)
    titles = [next(s for s in slide.shapes if s.name == "Title").text_frame.text for slide in prs.slides]
    assert titles == list(SLIDE_TITLES)
    for slide in prs.slides:  # the official logo, on the right of the header, aspect ratio kept
        logo = next(s for s in slide.shapes if s.name == "Logo")
        assert logo.left > prs.slide_width / 2 and logo.image.content_type == "image/png"
        width, height = logo.image.size
        assert abs(logo.width / logo.height - width / height) < 0.02
    text = pptx_text(path)
    assert all("Confidential" in t for t in text)
    assert "Placeholder industry services company.*" in text[0] and LEGEND_MANY in text[0]
    assert "A risk partner aligned to your workforce" in text[1] and "Why it matters:" in text[1]
    assert "— marsh.com — Retrieved" in text[1]  # slide 2's sources: the Marsh pages behind its statements
    assert any(s.has_table for s in prs.slides[2].shapes) and "₹2,50,000" in text[2]
    assert "Source: Niva Bupa ReAssure 2.0 Product Brochure, p. 2" in text[2]
    assert "Niva Bupa ReAssure 2.0" in text[3] and "Variant: Titanium+" in text[3]
    assert "Assumed sum insured: ₹10,00,000*" in text[3]
    assert f"[1] {SHARED_QUALIFIER}" in text[3] and SHARED_QUALIFIER not in text[2]  # on its own slide only
    assert "Frequent international travel makes this relevant.*" in text[3]  # framing sub-line
    for n in (0, 1):
        assert "the policy wording prevails" not in text[n]
    for n in (2, 3):
        assert "Summary based on insurer brochures; the policy wording prevails in case of conflict." in text[n]
    fonts = {r.font.name for slide in prs.slides for s in slide.shapes if s.has_text_frame
             for p in s.text_frame.paragraphs for r in p.runs}
    assert fonts == {FONT}


def test_nothing_internal_is_visible(rendered):
    _, path, _ = rendered
    joined = "\n".join(pptx_text(path))
    for word in (".md", ".pdf", "data/", "LLM", "Selected by", "Variant: not specified", "(Assumption)", "MS-",
                 "EV-", "audited", "MODEL_KNOWLEDGE"):
        assert word not in joined


def test_speaker_notes_trace_every_claim(rendered):
    ctx, path, _ = rendered
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(path).slides]
    assert all(notes)
    assert "CL-008 → VERIFIED → EV-NIVA-2-015 (Niva Bupa ReAssure 2.0, p. 2)" in notes[2]
    assert "CL-004 → VERIFIED → EV-MARSH-2-011 (Marsh profile, MS-021" in notes[1]
    assert "CL-001 → LABELLED_ASSUMPTION → CF-001 (company profile: Assumption)" in notes[0]
    assert "Recommended: Niva Bupa ReAssure 2.0 (POL-NIVA)" in notes[3]
    assert "Selected by:" in notes[3]  # how the policy was selected: notes and run data, never the slide


def test_one_assumption_gets_the_singular_legend():
    ctx = approved_run()
    ctx.deck.slides[0].bullets[1].state = ClaimState.REMOVED
    text = pptx_text(render(ctx)[0])
    assert LEGEND_ONE in text[0] and LEGEND_MANY not in text[0]


def test_slide4_hides_an_empty_variant_and_shows_add_ons():
    ctx = approved_run()
    ctx.deck.recommended.variant = None
    ctx.selection.selected_variant = None
    ctx.selected_documents[0].variants = []
    text = pptx_text(render(ctx)[0])[3]
    assert "Variant" not in text and "Add-on" not in text
    ctx.deck.recommended.required_addons = ["Placeholder add-on"]
    ctx.selection.required_addons = ["Placeholder add-on"]
    assert "Placeholder add-on" in pptx_text(render(ctx)[0])[3]


def test_only_passing_claims_are_rendered():
    ctx = approved_run()
    set_result(ctx, "CL-010", status=AuditStatus.UNSUPPORTED)  # not material: a review item, not rendered
    ctx.deck.slides[3].supporting_benefits[0].state = ClaimState.REMOVED
    set_result(ctx, "CL-011", status=AuditStatus.VERIFIED)
    set_result(ctx, "CL-009", status=AuditStatus.ADVISOR_ATTESTED, advisor_action=AdvisorAction.ATTESTED,
               advisor_note="Stated in the policy wording.")
    path, gate = render(ctx)
    text = pptx_text(path)
    assert gate.status == OverallFlag.REVIEW_REQUIRED
    assert "Frequent international travel" not in text[3] and "Road ambulance" not in text[3]
    assert f"{AIR} (Advisor-attested)" in text[3]
    notes = Presentation(path).slides[3].notes_slide.notes_text_frame.text
    assert "CL-010 → UNSUPPORTED → CF-005 (company profile: Assumption) [not rendered]" in notes


def test_render_refuses_on_fail():
    ctx = approved_run()
    set_result(ctx, "CL-011", status=AuditStatus.CONTRADICTED)
    with pytest.raises(RenderRefusedError, match="CONTRADICTED"):
        render(ctx)
    assert not (settings.OUTPUTS_DIR / ctx.run_id / "pitch.pptx").exists()


def test_a_web_sourced_fact_has_no_marker_and_its_page_as_the_source():
    ctx = approved_run()
    source = WebSource(source_id="WEB-001", url="https://www.example.com/about", title="About Example Co",
                       retrieved_at=datetime(2026, 9, 26, 10, 0).astimezone(), content="Placeholder page text.")
    ctx.company_profile.sources = [source]
    ctx.company_profile.facts[0].source_ids = ["WEB-001"]
    ctx.company_profile.facts[0].status = "WEB_SOURCED"
    first = ctx.deck.slides[0].bullets[0]
    first.text, first.qualifier_text = "Placeholder industry services company.", "Web-sourced"
    first.claim_type = ClaimType.COMPANY_FACT
    set_result(ctx, "CL-001", status=AuditStatus.VERIFIED)
    text = pptx_text(render(ctx)[0])[0]
    assert "Placeholder industry services company." in text and "company.*" not in text
    assert "About Example Co — example.com — Retrieved 26 September 2026" in text
    assert "Web-sourced" not in text and LEGEND_ONE in text  # the other bullet is still an assumption


def test_legacy_assumption_labels_become_markers():
    ctx = approved_run()
    ctx.deck.slides[0].bullets[0].text = "Placeholder industry services company. (Assumption)"
    text = pptx_text(render(ctx)[0])[0]
    assert "(Assumption)" not in text and "Placeholder industry services company.*" in text


# --- Structural QA ------------------------------------------------------------------------------------------------


def _broken(path, change) -> None:
    prs = Presentation(path)
    change(prs)
    prs.save(path)


def test_qa_detects_overflow(rendered):
    _, path, _ = rendered

    def shrink(prs):
        body = next(s for s in prs.slides[0].shapes if s.name == "Body")
        body.height = Emu(int(0.2 * 914400))

    _broken(path, shrink)
    with pytest.raises(RenderQAError, match="'Body' overflows"):
        structural_qa(path)


@pytest.mark.parametrize("change, message", [
    (lambda prs: setattr(next(s for s in prs.slides[0].shapes if s.name == "Title").text_frame, "text", "Intro"),
     "slide 1: title"),
    (lambda prs: setattr(next(s for s in prs.slides[1].shapes if s.name == "Headline").text_frame.paragraphs[0]
                         .runs[0], "text", "{{marsh}}"), "placeholder"),
    (lambda prs: setattr(next(s for s in prs.slides[0].shapes if s.name == "Body").text_frame.paragraphs[0].runs[0],
                         "text", "See data/policies/niva.pdf"), "'.pdf'"),
    (lambda prs: setattr(next(s for s in prs.slides[3].shapes if s.name == "RecommendedPolicy")
                         .text_frame.paragraphs[1].runs[0], "text", "Selected by: LLM"), "'Selected by'"),
    (lambda prs: setattr(next(s for s in prs.slides[0].shapes if s.name == "Body").text_frame.paragraphs[0].runs[0],
                         "text", "Large company (Assumption)"), r"'\(Assumption\)'"),
    (lambda prs: next(s for s in prs.slides[0].shapes if s.name == "SmallPrint").text_frame.clear(),
     "without the legend"),
    (lambda prs: setattr(next(s for s in prs.slides[2].shapes if s.name == "Logo"), "left", Emu(914400)),
     "logo"),
    (lambda prs: setattr(prs.slides[2].notes_slide.notes_text_frame, "text", ""), "slide 3: no speaker notes"),
    (lambda prs: setattr(next(s for s in prs.slides[3].shapes if s.name == "RecommendedPolicy")
                         .text_frame.paragraphs[1].runs[0], "text", "or HDFC ERGO Optima Secure+"), "exactly one"),
    (lambda prs: setattr(next(s for s in prs.slides[0].shapes if s.name == "Body").text_frame.paragraphs[0].runs[0],
                         "text", "Premium `22,616"), "backtick"),
])
def test_qa_detects_structural_problems(rendered, change, message):
    _, path, _ = rendered
    _broken(path, change)
    with pytest.raises(RenderQAError, match=message):
        structural_qa(path)


def test_qa_rejects_a_tiny_file(tmp_path):
    path = tmp_path / "pitch.pptx"
    path.write_bytes(b"x" * 100)
    with pytest.raises(RenderQAError, match="10 KB"):
        structural_qa(path)


def test_the_font_renders_the_rupee_sign():
    from marsh.render_ppt import _pil_font

    font = _pil_font(False, False)
    if font is None:
        pytest.skip("Arial is not installed here")
    assert bytes(font.getmask("₹")) != bytes(font.getmask(""))  # a real glyph, not the missing-glyph box
    assert text_width_pt("₹10,00,000", 12) > 0
